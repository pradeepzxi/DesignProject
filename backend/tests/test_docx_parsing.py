"""Unit Tests for Word DOCX Parsing with python-docx."""

from pathlib import Path
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from backend.app.services.parser_service import DocumentParserService


def add_hyperlink_test(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._element.append(hyperlink)


def test_docx_parsing_extracts_links_and_structure(tmp_path: Path):
    docx_path = tmp_path / "test_handbook.docx"

    doc = docx.Document()
    doc.core_properties.title = "Employee Handbook 2024"

    doc.add_heading("1.0 Code of Conduct", level=1)
    p = doc.add_paragraph("Please review the ")
    add_hyperlink_test(p, "https://example.com/ethics/charter", "Ethics Charter")
    p.add_run(" before starting work.")

    # Plain text link
    doc.add_paragraph("You can also view guidelines at https://example.com/plain/guidelines directly.")

    doc.save(docx_path)

    result = DocumentParserService.parse_docx(docx_path)

    assert result["document_title"] == "Employee Handbook 2024"
    assert len(result["links"]) >= 2

    ethics_link = next(l for l in result["links"] if "charter" in l["original_url"])
    assert ethics_link["original_url"] == "https://example.com/ethics/charter"
    assert ethics_link["anchor_text"] == "Ethics Charter"
    assert "Code of Conduct" in ethics_link["section_heading"]
