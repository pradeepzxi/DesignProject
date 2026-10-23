"""Tests for Autonomous Remediation, Backup creation, and Document Patching."""

import pytest
from pathlib import Path
from docx import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from backend.app.core.config import settings
from backend.app.services.remediation_service import AutonomousRemediationService
from backend.app.models.document import Document
from backend.app.models.link import Link
from backend.app.models.link_context import LinkContext


@pytest.mark.asyncio
async def test_docx_relationship_remediation_and_backup(db_session, tmp_path):
    # 1. Create a dummy docx file with a hyperlink
    sample_docx_path = tmp_path / "test_policy.docx"
    doc = DocxDocument()
    p = doc.add_paragraph("Please visit ")
    # Add relationship
    part = doc.part
    rel_id = part.relate_to("https://broken.corp/old-url", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    doc.save(str(sample_docx_path))

    # 2. Add records to test database
    db_doc = Document(
        filename="test_policy.docx",
        storage_path=str(sample_docx_path),
        file_type="docx",
        file_size=sample_docx_path.stat().st_size,
    )
    db_session.add(db_doc)
    db_session.commit()
    db_session.refresh(db_doc)

    db_link = Link(
        document_id=db_doc.document_id,
        original_url="https://broken.corp/old-url",
        normalized_url="https://broken.corp/old-url",
        status="BROKEN_404",
    )
    db_session.add(db_link)
    db_session.commit()
    db_session.refresh(db_link)

    # 3. Apply remediation
    service = AutonomousRemediationService(db_session)
    result = await service.apply_remediation(
        link_id=db_link.link_id,
        replacement_url="https://intranet.corp/new-verified-url",
        confidence_score=0.96,
        actor="SYSTEM",
        action="AUTO_REPAIRED",
    )

    assert result["success"] is True
    assert result["old_url"] == "https://broken.corp/old-url"
    assert result["new_url"] == "https://intranet.corp/new-verified-url"
    assert "backup_path" in result
    assert Path(result["backup_path"]).exists()

    # 4. Verify the patched docx actually contains the new URL in its relationships
    repaired_doc = DocxDocument(str(sample_docx_path))
    targets = [rel.target_ref for rel in repaired_doc.part.rels.values() if rel.reltype == RELATIONSHIP_TYPE.HYPERLINK]
    assert "https://intranet.corp/new-verified-url" in targets
    assert "https://broken.corp/old-url" not in targets
